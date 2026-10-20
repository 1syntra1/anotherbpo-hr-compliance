"""Calculation engines checked against independently hand-computed expectations,
plus programmatic inspection of the generated files."""
import io
import json
from datetime import date
from decimal import Decimal

import pytest

from conftest import query
from factories import add_employee, find_row, generate, open_record, versions, xlsx_values
from modules import bbbee, coida, ee
from modules.rule_seed import RULE_SETS

EE_PARAMS = next(r for r in RULE_SETS if r["code"] == "EE-ACT-2025")["params"]
GEN = next(r for r in RULE_SETS if r["code"] == "BBBEE-GENERIC-AMENDED")["params"]
QSE = next(r for r in RULE_SETS if r["code"] == "BBBEE-QSE-AMENDED")["params"]
EME = next(r for r in RULE_SETS if r["code"] == "BBBEE-EME-AMENDED")["params"]


def ee_snapshot(emps, start="2024-10-01", end="2025-09-30"):
    base = {"id": 0, "emp_number": None, "race": "African", "gender": "Male", "disability": "No", "foreign": False,
            "level": "Unskilled", "employment_type": "Permanent", "start": "2020-01-01", "end": None,
            "status": "ACTIVE", "monthly_salary": 10000}
    rows = []
    for i, e in enumerate(emps, 1):
        rows.append(dict(base, id=i, emp_number="X%d" % i, **e))
    return {"kind": "ee", "period": {"label": "t", "start": start, "end": end}, "company": {},
            "rules": {"id": 1, "code": "EE", "version": "t", "verified": False, "params": EE_PARAMS},
            "employees": rows}


# ── Employment Equity ──────────────────────────────────────────

def test_ee_zero_employees():
    a = ee.analyse(ee_snapshot([]))
    checks = ee.reconcile(a, ee.eea2_data(a), ee.eea4_data(a))
    assert a["population"] == [] and all(c["ok"] for c in checks)


def test_ee_one_employee_full_year():
    a = ee.analyse(ee_snapshot([{"monthly_salary": 12345.67}]))
    assert len(a["population"]) == 1
    # Full period (365 days): 12 months exactly
    assert Decimal(a["population"][0]["remuneration"]) == Decimal("148148.04")


def test_ee_population_rules():
    emps = [
        {},                                                     # 1 full year
        {"start": "2025-03-01"},                                # 2 hired in period
        {"end": "2025-06-30", "status": "TERMINATED"},          # 3 terminated in period
        {"start": "2025-10-01"},                                # 4 hired after period
        {"end": "2024-09-30", "status": "TERMINATED"},          # 5 left before period
        {"foreign": True, "race": None, "gender": "Female"},    # 6 foreign national
        {"employment_type": "Temporary"},                       # 7 temporary
        {"gender": None},                                       # 8 missing gender → excluded
        {"end": "2025-09-30", "status": "TERMINATED"},          # 9 last day = still employed at end
    ]
    a = ee.analyse(ee_snapshot(emps))
    pop = {e["emp_number"] for e in a["population"]}
    assert pop == {"X1", "X2", "X6", "X7", "X9"}
    assert {e["emp_number"] for e in a["hires"]} == {"X2"}
    assert {e["emp_number"] for e in a["terms"]} == {"X3", "X9"}
    assert [e["emp_number"] for e in a["excluded"]] == ["X8"]
    e2 = ee.eea2_data(a)
    assert e2["permanent"]["Unskilled"][("Foreign", "F")] == 1
    assert e2["temporary"]["Unskilled"][("African", "M")] == 1
    checks = ee.reconcile(a, e2, ee.eea4_data(a))
    failed = [c["check"] for c in checks if not c["ok"]]
    assert failed == ["No employees excluded by data errors"]


def test_ee_part_year_remuneration_prorated_by_days():
    # Hired 1 April 2025 → 183 of 365 days in period
    a = ee.analyse(ee_snapshot([{"start": "2025-04-01", "monthly_salary": 10000}]))
    expected = (Decimal(10000) * (Decimal(183) / Decimal(365) * 12).quantize(Decimal("0.0001"))).quantize(Decimal("0.01"))
    assert Decimal(a["population"][0]["remuneration"]) == expected == Decimal("60164.00")


@pytest.mark.parametrize("salary,severity", [(None, "error"), (-1, "error"), (0, "warning")])
def test_ee_bad_salaries_flagged(salary, severity):
    a = ee.analyse(ee_snapshot([{"monthly_salary": salary}]))
    assert any(i["severity"] == severity for i in a["issues"])


def test_ee_high_salary_and_rounding():
    a = ee.analyse(ee_snapshot([{"monthly_salary": 999999.995}, {"monthly_salary": 0.005}]))
    total = sum(Decimal(e["remuneration"]) for e in a["population"])
    e4 = ee.eea4_data(a)
    assert sum(v for r in e4["remuneration"].values() for v in r.values()) == total


def test_ee_reconciliation_detects_tampered_eea4():
    a = ee.analyse(ee_snapshot([{}, {"gender": "Female"}]))
    e2, e4 = ee.eea2_data(a), ee.eea4_data(a)
    e4["headcount"]["Unskilled"][("African", "F")] -= 1   # simulate a discrepancy
    bad = [c for c in ee.reconcile(a, e2, e4) if not c["ok"]]
    names = " ".join(c["check"] for c in bad)
    assert "EEA2 total = EEA4 total" in names and "Unskilled" in names and "African F" in names


def test_ee_no_period_dates():
    a = ee.analyse(ee_snapshot([{}], start=None, end=None))
    assert a["issues"][0]["severity"] == "error"


# ── Generated Excel inspection ─────────────────────────────────

ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A")


def _inspect(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", "formula in %s!%s" % (ws.title, c.coordinate)
                if isinstance(c.value, str):
                    assert not any(e in c.value for e in ERRORS)
    return wb


@pytest.mark.parametrize("n", [0, 1, 10, 100, 1000])
def test_eea_reports_scale_and_reconcile(app_module, admin_client, n):
    db = app_module.connect_db()
    levels = EE_PARAMS["occupational_levels"]
    races = EE_PARAMS["race_groups"]
    for i in range(n):
        db.execute("""INSERT INTO employees (emp_number, first_name, last_name, race, gender, occupational_level,
                      employment_type, salary, employment_date, status) VALUES (?,?,?,?,?,?,?,?,?, 'ACTIVE')""",
                   ("S%05d" % i, "F", "L", races[i % 4], "Male" if i % 3 else "Female", levels[i % 6],
                    "Temporary" if i % 10 == 0 else "Permanent", 8000 + i, "2019-01-01"))
    db.commit(); db.close()
    rid = open_record(admin_client, "ee", 2025)
    generate(admin_client, rid, "eea2"); generate(admin_client, rid, "eea4")
    vs = {v["report_type"]: v for v in versions(app_module, rid)}
    assert vs["eea2"]["checks_ok"] and vs["eea4"]["checks_ok"]
    d2 = admin_client.get("/reports/%d/download" % vs["eea2"]["id"])
    assert d2.headers["Content-Type"].startswith("application/vnd.openxmlformats")
    assert d2.headers["Content-Disposition"].endswith("v1.xlsx")
    wb = _inspect(d2.data)
    assert wb.sheetnames == ["EEA2", "Checks"]
    rows = xlsx_values(d2.data)["EEA2"]
    totals = [r for r in rows if r and r[0] == "TOTAL"]
    perm = sum(1 for i in range(n) if i % 10)
    assert totals[0][-1] == perm and totals[1][-1] == n - perm
    wb4 = _inspect(admin_client.get("/reports/%d/download" % vs["eea4"]["id"]).data)
    assert wb4.sheetnames == ["EEA4", "Checks"]
    rows4 = [[c.value for c in r] for r in wb4["EEA4"].iter_rows()]
    rem_total = [r for r in rows4 if r and r[0] == "TOTAL"][1][-1]
    expected = float(sum(Decimal(8000 + i) * 12 for i in range(n)))
    assert abs(rem_total - expected) < 0.005


# ── COIDA ──────────────────────────────────────────────────────

def coida_snap(earnings, rates, year="2024"):
    return {"kind": "coida", "period": {"label": year, "start": None, "end": None}, "company": {},
            "rules": {"id": 1, "code": "C", "version": "t", "verified": False,
                      "params": {"max_earnings_by_year": {"2024": 563520.0}}},
            "earnings": [dict(employee_id=i, employee_number="E%d" % i, employee_name="n", months_worked=12,
                              subclass="", **e) for i, e in enumerate(earnings)],
            "rates": [dict(subclass="", description="", source_note="", **r) for r in rates]}


def test_coida_assessment_formula():
    s = coida_snap([{"class_code": "A", "earnings": 100000}, {"class_code": "A", "earnings": 250000.55},
                    {"class_code": "B", "earnings": 600000}],          # B is above the cap
                   [{"class_code": "A", "rate": 0.48}, {"class_code": "B", "rate": 1.23}])
    c = coida.calculate(s)
    a, b = c["classes"]
    assert a["capped"] == Decimal("350000.55")
    assert a["assessment"] == (Decimal("350000.55") / 100 * Decimal("0.48")).quantize(Decimal("0.01"))  # 1680.00
    assert b["capped"] == Decimal("563520.00")
    assert b["assessment"] == Decimal("6931.30")                         # 563520 / 100 * 1.23
    assert c["total_assessment"] == a["assessment"] + b["assessment"]
    assert all(ch["ok"] for ch in coida.checks(s, c))


def test_coida_missing_rate_and_unknown_year():
    s = coida_snap([{"class_code": "Z", "earnings": 1000}], [], year="2031")
    c = coida.calculate(s)
    failed = {ch["check"] for ch in coida.checks(s, c) if not ch["ok"]}
    assert failed == {"Every class has an assessment rate", "Maximum earnings known for the period"}


@pytest.mark.parametrize("earn", [0, 0.01, 99.995])
def test_coida_small_values(earn):
    s = coida_snap([{"class_code": "A", "earnings": earn}], [{"class_code": "A", "rate": 1.0}])
    c = coida.calculate(s)
    assert c["total_assessment"] == (coida.money(earn) / 100).quantize(Decimal("0.01"))


def test_coida_negative_earnings_rejected_by_database(app_module):
    import sqlite3
    db = app_module.connect_db()
    rid = db.execute("SELECT id FROM compliance_records WHERE module='coida'").fetchone()[0]
    with pytest.raises(sqlite3.IntegrityError):
        db.execute("INSERT INTO coida_earnings (record_id, class_code, earnings) VALUES (?, 'A', -1)", (rid,))
    db.close()


def test_coida_end_to_end_build_and_report(app_module, admin_client):
    from factories import form_post
    add_employee(app_module, "C1", salary=10000, coida_class_code="A")
    add_employee(app_module, "C2", salary=20000, coida_class_code="A", employment_date="2024-09-01")
    add_employee(app_module, "C3", salary=50000, coida_class_code="B")
    add_employee(app_module, "C4", salary=5000, coida_class_code=None)           # skipped
    rid = open_record(admin_client, "coida", 2024)                              # Mar 2024 – Feb 2025
    form_post(admin_client, "/coida/build-earnings", {"record_id": rid}, "/coida?record=%d" % rid)
    for code, rate in (("A", "0.48"), ("B", "1.23")):
        form_post(admin_client, "/coida/rate", {"record_id": rid, "class_code": code, "rate": rate}, "/coida")
    rows = {r["employee_number"]: r for r in query(app_module, "SELECT * FROM coida_earnings WHERE record_id=?", (rid,))}
    assert set(rows) == {"C1", "C2", "C3"}
    assert rows["C1"]["earnings"] == 120000.0
    # C2: 1 Sep 2024 – 28 Feb 2025 = 181 of 365 days
    assert rows["C2"]["earnings"] == float((Decimal(20000) * (Decimal(181) / 365 * 12).quantize(Decimal("0.0001"))).quantize(Decimal("0.01")))
    generate(admin_client, rid, "was8")
    v = versions(app_module, rid, "was8")[0]
    assert v["checks_ok"]
    data = admin_client.get("/reports/%d/download" % v["id"]).data
    _inspect(data)
    vals = xlsx_values(data)
    assert set(vals) == {"W.As.8", "Employees", "Checks"}
    total = find_row(vals["W.As.8"], "TOTAL")
    exp_a = ((Decimal("120000") + Decimal(str(rows["C2"]["earnings"]))) / 100 * Decimal("0.48")).quantize(Decimal("0.01"))
    exp_b = (Decimal("563520") / 100 * Decimal("1.23")).quantize(Decimal("0.01"))
    assert abs(total[-1] - float(exp_a + exp_b)) < 0.005


# ── B-BBEE ─────────────────────────────────────────────────────

FULL_SUBMIN = {"ownership_net_value": 8, "skills_total": 20, "esd_preferential_procurement": 25,
               "esd_supplier_development": 10, "esd_enterprise_development": 5}


def generic_inputs(extra=0):
    """All priority elements at maximum with sub-minimums met (85 points), plus `extra`
    points spread over management, SED and bonuses."""
    out = {"ownership": "25", "skills": "20", "esd": "40", "management": "0", "sed": "0"}
    left = Decimal(str(extra))
    for k, cap in (("management", 19), ("sed", 5), ("skills_bonus", 5), ("esd_bonus", 4)):
        take = min(left, Decimal(cap))
        out[k] = str(take)
        left -= take
    out.update({k: str(v) for k, v in FULL_SUBMIN.items()})
    return out


@pytest.mark.parametrize("points,level", [
    (113, 1), (100, 1), (99.99, 2), (95, 2), (94.99, 3), (90, 3), (89.99, 4), (80, 4), (79.99, 5), (75, 5),
    (74.99, 6), (70, 6), (69.99, 7), (55, 7), (54.99, 8), (40, 8), (39.99, "NC"), (0, "NC"),
])
def test_bbbee_level_boundaries(points, level):
    assert bbbee.level_for_points(Decimal(str(points)), GEN["level_thresholds"]) == level
    assert bbbee.level_for_points(Decimal(str(points)), QSE["level_thresholds"]) == level


@pytest.mark.parametrize("extra,total,level", [(0, 85, 4), (5, 90, 3), (15, 100, 1), (28, 113, 1), (4.99, 89.99, 4)])
def test_bbbee_generic_full_calculation(extra, total, level):
    r = bbbee.calculate("GENERIC", GEN, generic_inputs(extra))
    assert r["total"] == Decimal(str(total)).quantize(Decimal("0.01")) and r["level"] == level
    assert not r["discounted"] and r["recognition_pct"] == GEN["recognition_pct"][str(level)]


def test_bbbee_discounting_one_level():
    inp = generic_inputs(11)                          # 96 points → Level 2 before discounting
    inp["skills_total"] = "7.99"                      # 40% of 20 = 8 → not met
    r = bbbee.calculate("GENERIC", GEN, inp)
    assert r["base_level"] == 2 and r["level"] == 3 and r["discounted"]
    assert any("Discounted by 1 level" in n and "Skills Development" in n for n in r["notes"])


def test_bbbee_discount_from_level_8_is_non_compliant():
    inp = {"ownership": "10", "management": "10", "skills": "10", "esd": "10", "sed": "5", "ownership_net_value": "1",
           "skills_total": "10", "esd_preferential_procurement": "10", "esd_supplier_development": "4",
           "esd_enterprise_development": "2"}
    r = bbbee.calculate("GENERIC", GEN, inp)
    assert r["base_level"] == 8 and r["level"] == "NC"


def test_bbbee_qse_priority_rule_ownership_plus_one():
    inp = {"ownership": "20", "management": "15", "skills": "25", "esd": "30", "sed": "5",
           "ownership_net_value": "8", "skills_total": "25", "esd_preferential_procurement": "0",
           "esd_supplier_development": "0", "esd_enterprise_development": "0"}
    r = bbbee.calculate("QSE", QSE, inp)
    assert not r["discounted"]                       # ownership + skills met; ESD failing is allowed
    inp["skills_total"] = "0"
    r = bbbee.calculate("QSE", QSE, inp)
    assert r["discounted"]


@pytest.mark.parametrize("bo,level", [(100, 1), (51, 2), (50.99, 4), (0, 4)])
def test_bbbee_eme_automatic_levels(bo, level):
    assert bbbee.calculate("EME", EME, {}, bo)["level"] == level


def test_bbbee_qse_black_owned_automatic_level_wins():
    r = bbbee.calculate("QSE", QSE, {k: "0" for k in ("ownership_net_value",)}, 100)
    assert r["level"] == 1


@pytest.mark.parametrize("bad", [{"ownership": "26"}, {"ownership": "-1"}, {"skills_bonus": "6"},
                                 {"ownership": "abc"}, {"ownership_net_value": "9"}])
def test_bbbee_invalid_inputs(bad):
    inp = generic_inputs(0)
    inp.update(bad)
    with pytest.raises(bbbee.InputError):
        bbbee.calculate("GENERIC", GEN, inp)


def test_bbbee_classification():
    assert bbbee.classify(10_000_000, GEN) == "EME"
    assert bbbee.classify(10_000_000.01, GEN) == "QSE"
    assert bbbee.classify(50_000_000.01, GEN) == "GENERIC"


def test_bbbee_reports_never_claim_to_be_certificates(app_module, admin_client):
    from factories import form_post
    rid = query(app_module, "SELECT id FROM compliance_records WHERE module='bbbee'")[0][0]
    data = dict(generic_inputs(0), record_id=rid, scorecard_type="GENERIC", annual_turnover="60000000",
                black_owned_pct="30", black_women_owned_pct="10")
    form_post(admin_client, "/bbbee/save", data, "/bbbee")
    generate(admin_client, rid, "bbbee_pdf"); generate(admin_client, rid, "bbbee_xlsx")
    vs = {v["report_type"]: v for v in versions(app_module, rid)}
    pdf = admin_client.get("/reports/%d/download" % vs["bbbee_pdf"]["id"])
    assert pdf.data.startswith(b"%PDF") and pdf.headers["Content-Type"] == "application/pdf"
    assert "Certificate" not in vs["bbbee_pdf"]["file_name"]
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf.data))
    text = " ".join(p.extract_text() for p in reader.pages)
    assert len(reader.pages) >= 1 and reader.metadata.title == "Internal B-BBEE Calculation"
    assert "INTERNAL B-BBEE CALCULATION" in text and "NOT a B-BBEE verification certificate" in text
    assert "Level 4" in text
    vals = xlsx_values(admin_client.get("/reports/%d/download" % vs["bbbee_xlsx"]["id"]).data)
    flat = " ".join(str(c) for row in vals["Scorecard"] for c in row if c)
    assert "NOT a B-BBEE certificate" in flat and "Level 4" in flat


# ── WSP / ATR ──────────────────────────────────────────────────

def test_wsp_reports_and_word_structure(app_module, admin_client):
    from factories import form_post
    eid = add_employee(app_module, "W1", first_name="Zoë", last_name="O'Brien-Nkosi")
    rid = query(app_module, "SELECT id FROM compliance_records WHERE module='wsp'")[0][0]
    for status_, actual in (("Planned", ""), ("Completed", "2026-02-01")):
        form_post(admin_client, "/wsp/training", {"record_id": rid, "employee_id": eid, "intervention_type": "Excel",
                                                  "cost": "1500.50", "status": status_, "actual_date": actual}, "/wsp")
    # employee later archived: training history keeps the name/number
    db = app_module.connect_db(); db.execute("UPDATE employees SET status='ARCHIVED', first_name='Changed' WHERE id=?", (eid,)); db.commit(); db.close()
    for t in ("wsp", "atr", "wsp_docx"):
        generate(admin_client, rid, t)
    vs = {v["report_type"]: v for v in versions(app_module, rid)}
    assert all(v["checks_ok"] for v in vs.values())
    wsp_vals = xlsx_values(admin_client.get("/reports/%d/download" % vs["wsp"]["id"]).data)
    assert find_row(wsp_vals["WSP"], 1)[2] == "Zoë O'Brien-Nkosi"
    assert find_row(wsp_vals["WSP"], "TOTAL")[6] == 1500.5
    from docx import Document
    doc = Document(io.BytesIO(admin_client.get("/reports/%d/download" % vs["wsp_docx"]["id"]).data))
    assert len(doc.tables) == 3 and doc.tables[1].rows[1].cells[1].text == "Zoë O'Brien-Nkosi"


@pytest.mark.parametrize("field", ["intervention_type", "provider"])
def test_formula_injection_refused_at_input(app_module, admin_client, field):
    from factories import form_post
    eid = add_employee(app_module, "W2")
    rid = query(app_module, "SELECT id FROM compliance_records WHERE module='wsp'")[0][0]
    data = {"record_id": rid, "employee_id": eid, "intervention_type": "Course", "status": "Planned"}
    data[field] = '=HYPERLINK("http://evil.example","x")'
    form_post(admin_client, "/wsp/training", data, "/wsp")
    assert not query(app_module, "SELECT 1 FROM training_records")


def test_formula_neutralised_in_output():
    from openpyxl import Workbook
    from modules.xlsx_safe import neutralise_workbook
    wb = Workbook()
    wb.active["A1"] = '=cmd|"/c calc"!A1'
    assert wb.active["A1"].data_type == "f"
    neutralise_workbook(wb)
    assert wb.active["A1"].data_type == "s"
