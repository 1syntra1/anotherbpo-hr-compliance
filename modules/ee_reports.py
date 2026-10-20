"""EEA2 / EEA4 Excel output, rendered purely from a snapshot (see modules/ee.py).

These workbooks are internal working reports that mirror the content of the
statutory forms; the official submission is made on the Department's online
system. Layout should be checked against the current forms (rule set notes).
"""
from decimal import Decimal

from openpyxl.styles import Font

from modules import ee
from modules.xlsx_common import (data_row, header_row, new_workbook, report_meta, save, title_block,
                                 warning_row, widths)

LEVEL_SHORT = {
    "Top Management": "Top management",
    "Senior Management": "Senior management",
    "Professionally Qualified": "Professionally qualified & experienced specialists / mid-management",
    "Skilled Technical": "Skilled technical & academically qualified / junior management",
    "Semi-Skilled": "Semi-skilled & discretionary decision making",
    "Unskilled": "Unskilled & defined decision making",
}


def _col_headers(cols):
    out = []
    for grp, g in cols:
        out.append("%s %s" % ("Foreign" if grp == "Foreign" else grp, "Male" if g == "M" else "Female"))
    return out + ["Total"]


def _company_meta(s):
    c = s["company"]
    p = s["period"]
    return [
        ("Employer:", c.get("company_name") or ""),
        ("Registration number:", c.get("reg_number") or ""),
        ("Designated employer:", c.get("designated_employer") or ""),
        ("Sector:", c.get("sector") or ""),
        ("Reporting period:", "%s  (%s to %s)" % (p["label"], p["start"], p["end"])),
    ]


def _matrix_block(ws, row, title, m, levels, cols):
    ws.cell(row=row, column=1, value=title).font = Font(bold=True)
    row += 1
    row = header_row(ws, row, ["Occupational level"] + _col_headers(cols))
    totals = [0] * (len(cols) + 1)
    for i, lvl in enumerate(levels):
        vals = [m[lvl][c] for c in cols]
        vals.append(sum(vals))
        totals = [a + b for a, b in zip(totals, vals)]
        row = data_row(ws, row, [LEVEL_SHORT.get(lvl, lvl)] + vals, shade=i % 2 == 0)
    row = data_row(ws, row, ["TOTAL"] + totals, total=True)
    return row + 1


def _issues_sheet(wb, a, checks):
    ws = wb.create_sheet("Checks")
    widths(ws, [52, 18, 18, 8, 60])
    row = header_row(ws, 1, ["Check", "Expected", "Actual", "OK", "Detail"])
    for c in checks:
        row = data_row(ws, row, [c["check"], c["expected"], c["actual"], "PASS" if c["ok"] else "FAIL",
                                 c["detail"]])
    row += 1
    row = header_row(ws, row, ["Data issue", "Employee", "Severity", "", ""])
    for i in a["issues"]:
        row = data_row(ws, row, [i["message"], i["employee"] or "", i["severity"], "", ""])


def generate_eea2(snapshot, ctx):
    a = ee.analyse(snapshot)
    e2 = ee.eea2_data(a) if a.get("levels") else None
    cols = ee.columns(snapshot["rules"]["params"]["race_groups"])
    width = len(cols) + 2
    wb, ws = new_workbook("EEA2")
    widths(ws, [48] + [11] * (width - 1))
    row = title_block(ws, "EEA2 — EMPLOYMENT EQUITY REPORT (WORKING COPY)",
                      "Employment Equity Act 55 of 1998 — workforce profile and movements", width,
                      _company_meta(snapshot) + report_meta(ctx))
    if not snapshot["rules"]["verified"]:
        row = warning_row(ws, row, "Rule set %s has not been verified against the official source. Check the "
                                   "figures and the form layout before submission." % snapshot["rules"]["code"], width)
    if e2 is None:
        row = warning_row(ws, row, "No figures: the reporting period has no dates.", width)
    else:
        lv = a["levels"]
        row = _matrix_block(ws, row, "1. Workforce profile — permanent employees", e2["permanent"], lv, cols)
        row = _matrix_block(ws, row, "2. Workforce profile — temporary employees", e2["temporary"], lv, cols)
        row = _matrix_block(ws, row, "3. Employees with disabilities (permanent and temporary)", e2["disabled"],
                            lv, cols)
        row = _matrix_block(ws, row, "4. Recruitment during the period", e2["recruitment"], lv, cols)
        row = _matrix_block(ws, row, "5. Terminations during the period", e2["terminations"], lv, cols)
    checks = ctx.get("checks") or (ee.reconcile(a, e2, ee.eea4_data(a)) if e2 else [])
    _issues_sheet(wb, a, checks)
    return save(wb, ctx.get("generated_dt"))


def generate_eea4(snapshot, ctx):
    a = ee.analyse(snapshot)
    cols = ee.columns(snapshot["rules"]["params"]["race_groups"])
    width = len(cols) + 2
    wb, ws = new_workbook("EEA4")
    widths(ws, [48] + [14] * (width - 1))
    row = title_block(ws, "EEA4 — INCOME DIFFERENTIAL STATEMENT (WORKING COPY)",
                      "Remuneration for the reporting period, by occupational level, race and gender", width,
                      _company_meta(snapshot) + report_meta(ctx) +
                      [("Remuneration basis:", "Monthly salary x months employed in period (pro-rated)")])
    if not snapshot["rules"]["verified"]:
        row = warning_row(ws, row, "Rule set %s has not been verified against the official source."
                          % snapshot["rules"]["code"], width)
    if not a.get("levels"):
        row = warning_row(ws, row, "No figures: the reporting period has no dates.", width)
        _issues_sheet(wb, a, ctx.get("checks") or [])
        return save(wb, ctx.get("generated_dt"))
    e4 = ee.eea4_data(a)
    lv = a["levels"]
    ws.cell(row=row, column=1, value="A. Headcount")
    row = header_row(ws, row + 1, ["Occupational level"] + _col_headers(cols))
    tot = [0] * (len(cols) + 1)
    for i, lvl in enumerate(lv):
        vals = [e4["headcount"][lvl][c] for c in cols]
        vals.append(sum(vals))
        tot = [x + y for x, y in zip(tot, vals)]
        row = data_row(ws, row, [LEVEL_SHORT.get(lvl, lvl)] + vals, shade=i % 2 == 0)
    row = data_row(ws, row, ["TOTAL"] + tot, total=True) + 1
    ws.cell(row=row, column=1, value="B. Total remuneration (R)")
    row = header_row(ws, row + 1, ["Occupational level"] + _col_headers(cols))
    money_cols = range(2, width + 1)
    tot = [Decimal("0.00")] * (len(cols) + 1)
    for i, lvl in enumerate(lv):
        vals = [e4["remuneration"][lvl][c] for c in cols]
        vals.append(sum(vals, Decimal("0.00")))
        tot = [x + y for x, y in zip(tot, vals)]
        row = data_row(ws, row, [LEVEL_SHORT.get(lvl, lvl)] + [float(v) for v in vals],
                       money_cols=money_cols, shade=i % 2 == 0)
    row = data_row(ws, row, ["TOTAL"] + [float(v) for v in tot], money_cols=money_cols, total=True) + 1
    ws.cell(row=row, column=1, value="C. Average remuneration per employee (R)")
    row = header_row(ws, row + 1, ["Occupational level"] + _col_headers(cols)[:-1])
    for i, lvl in enumerate(lv):
        vals = []
        for c in cols:
            n = e4["headcount"][lvl][c]
            vals.append(float(ee.money(e4["remuneration"][lvl][c] / n)) if n else 0.0)
        row = data_row(ws, row, [LEVEL_SHORT.get(lvl, lvl)] + vals, money_cols=money_cols, shade=i % 2 == 0)
    checks = ctx.get("checks") or ee.reconcile(a, ee.eea2_data(a), e4)
    _issues_sheet(wb, a, checks)
    return save(wb, ctx.get("generated_dt"))
